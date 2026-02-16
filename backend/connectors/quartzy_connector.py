"""
Quartzy Connector
Connects to Quartzy REST API to import lab inventory items and order requests.
Also supports direct CSV/Excel file upload from Quartzy exports.
"""

import os
import csv
import io
import time
from datetime import datetime, timezone
from typing import List, Dict, Optional, Any

from connectors.base_connector import BaseConnector, ConnectorConfig, ConnectorStatus, Document

try:
    import requests as http_requests
    REQUESTS_AVAILABLE = True
except ImportError:
    REQUESTS_AVAILABLE = False


class QuartzyConnector(BaseConnector):
    """
    Quartzy connector for importing lab inventory and order requests.

    Two modes:
    1. API Access Token: Fetches inventory items + order requests via REST API
    2. CSV Upload: Parses Quartzy CSV/Excel exports (handled by static method)

    API Auth: Access-Token header (generated in Quartzy user settings)
    API Base: https://api.quartzy.com
    """

    CONNECTOR_TYPE = "quartzy"
    REQUIRED_CREDENTIALS = ["access_token"]
    OPTIONAL_SETTINGS = {
        "max_items": 500,
        "include_order_requests": True,
    }

    BASE_URL = "https://api.quartzy.com"

    def __init__(self, config: ConnectorConfig):
        super().__init__(config)
        self.access_token = config.credentials.get("access_token", "")
        self.headers = {
            "Access-Token": self.access_token,
            "Accept": "application/json"
        }

    async def connect(self) -> bool:
        """Validate access token against Quartzy API"""
        if not REQUESTS_AVAILABLE:
            self._set_error("requests library not installed")
            return False

        if not self.access_token:
            self._set_error("Missing access_token")
            return False

        try:
            self.status = ConnectorStatus.CONNECTING

            # Validate token by fetching first page of order requests
            response = http_requests.get(
                f"{self.BASE_URL}/order-requests",
                headers=self.headers,
                params={"page": 1},
                timeout=15
            )

            if response.status_code == 401:
                self._set_error("Invalid access token. Generate one in Quartzy Settings > API.")
                return False

            if response.status_code != 200:
                self._set_error(f"Quartzy API returned status {response.status_code}")
                return False

            self.status = ConnectorStatus.CONNECTED
            self._clear_error()
            print("[Quartzy] Connected successfully")
            return True

        except Exception as e:
            self._set_error(f"Connection failed: {str(e)}")
            return False

    async def disconnect(self) -> bool:
        """Disconnect from Quartzy"""
        self.status = ConnectorStatus.DISCONNECTED
        return True

    async def test_connection(self) -> bool:
        """Test if connection is valid"""
        return await self.connect()

    async def sync(self, since: Optional[datetime] = None) -> List[Document]:
        """
        Fetch inventory items and order requests from Quartzy API.

        Returns list of Document objects ready for embedding.
        """
        if self.status != ConnectorStatus.CONNECTED:
            if not await self.connect():
                return []

        self.status = ConnectorStatus.SYNCING
        documents = []

        try:
            max_items = self.config.settings.get("max_items", 500)
            include_orders = self.config.settings.get("include_order_requests", True)

            print(f"[Quartzy] Starting sync (max_items={max_items})")

            # Fetch inventory items
            inventory_docs = self._fetch_inventory_items(max_items)
            documents.extend(inventory_docs)
            print(f"[Quartzy] Fetched {len(inventory_docs)} inventory items")

            # Fetch order requests
            if include_orders:
                order_docs = self._fetch_order_requests(max_items)
                documents.extend(order_docs)
                print(f"[Quartzy] Fetched {len(order_docs)} order requests")

            self.config.last_sync = datetime.now(timezone.utc)
            self.status = ConnectorStatus.CONNECTED
            self._clear_error()
            print(f"[Quartzy] Sync complete: {len(documents)} total documents")

        except Exception as e:
            self._set_error(f"Sync failed: {str(e)}")
            print(f"[Quartzy] Sync error: {e}")
            import traceback
            traceback.print_exc()

        return documents

    async def get_document(self, doc_id: str) -> Optional[Document]:
        """Get a specific document by ID"""
        return None

    def _fetch_inventory_items(self, max_items: int) -> List[Document]:
        """Fetch inventory items via paginated API"""
        documents = []
        page = 1
        per_page = 50

        while len(documents) < max_items:
            try:
                time.sleep(0.5)  # Rate limiting
                response = http_requests.get(
                    f"{self.BASE_URL}/inventory-items",
                    headers=self.headers,
                    params={"page": page, "per_page": per_page},
                    timeout=30
                )

                if response.status_code != 200:
                    print(f"[Quartzy] Inventory API error (page {page}): {response.status_code}")
                    break

                data = response.json()
                items = data.get("data", data.get("inventory_items", data if isinstance(data, list) else []))

                if not items:
                    break

                for item in items:
                    if len(documents) >= max_items:
                        break

                    doc = self._inventory_item_to_document(item)
                    if doc:
                        documents.append(doc)

                # Check if more pages
                meta = data.get("meta", {})
                total_pages = meta.get("total_pages", meta.get("last_page", page))
                if page >= total_pages:
                    break

                page += 1

            except Exception as e:
                print(f"[Quartzy] Error fetching inventory page {page}: {e}")
                break

        return documents

    def _fetch_order_requests(self, max_items: int) -> List[Document]:
        """Fetch order requests via paginated API"""
        documents = []
        page = 1
        per_page = 50

        while len(documents) < max_items:
            try:
                time.sleep(0.5)  # Rate limiting
                response = http_requests.get(
                    f"{self.BASE_URL}/order-requests",
                    headers=self.headers,
                    params={"page": page, "per_page": per_page},
                    timeout=30
                )

                if response.status_code != 200:
                    print(f"[Quartzy] Order API error (page {page}): {response.status_code}")
                    break

                data = response.json()
                requests_list = data.get("data", data.get("order_requests", data if isinstance(data, list) else []))

                if not requests_list:
                    break

                for req in requests_list:
                    if len(documents) >= max_items:
                        break

                    doc = self._order_request_to_document(req)
                    if doc:
                        documents.append(doc)

                # Check if more pages
                meta = data.get("meta", {})
                total_pages = meta.get("total_pages", meta.get("last_page", page))
                if page >= total_pages:
                    break

                page += 1

            except Exception as e:
                print(f"[Quartzy] Error fetching orders page {page}: {e}")
                break

        return documents

    def _inventory_item_to_document(self, item: Dict[str, Any]) -> Optional[Document]:
        """Convert a Quartzy inventory item to Document"""
        try:
            item_id = str(item.get("id", item.get("serial_number", "")))
            name = item.get("name", item.get("item_name", "Untitled Item"))

            # Build rich text content for embedding
            parts = [f"Lab Inventory Item: {name}"]

            if item.get("vendor"):
                parts.append(f"Vendor: {item['vendor']}")
            if item.get("catalog_number") or item.get("catalog_#"):
                parts.append(f"Catalog #: {item.get('catalog_number', item.get('catalog_#', ''))}")
            if item.get("location"):
                parts.append(f"Location: {item['location']}")
            if item.get("owner"):
                parts.append(f"Owner: {item['owner']}")
            if item.get("unit_size"):
                parts.append(f"Unit Size: {item['unit_size']}")
            if item.get("amount_in_stock") is not None:
                parts.append(f"Amount in Stock: {item['amount_in_stock']}")
            if item.get("type"):
                parts.append(f"Type: {item['type']}")
            if item.get("description"):
                parts.append(f"Description: {item['description']}")
            if item.get("notes"):
                parts.append(f"Notes: {item['notes']}")
            if item.get("expiration_date"):
                parts.append(f"Expiration Date: {item['expiration_date']}")
            if item.get("lot_number"):
                parts.append(f"Lot Number: {item['lot_number']}")

            content = "\n".join(parts)

            timestamp = None
            if item.get("updated_at") or item.get("created_at"):
                try:
                    ts = item.get("updated_at") or item.get("created_at")
                    timestamp = datetime.fromisoformat(str(ts).replace("Z", "+00:00"))
                except Exception:
                    pass

            return Document(
                doc_id=f"quartzy_inv_{item_id}",
                source="quartzy",
                content=content,
                title=f"Inventory: {name}",
                metadata={
                    "quartzy_id": item_id,
                    "item_type": "inventory",
                    "vendor": item.get("vendor", ""),
                    "catalog_number": item.get("catalog_number", ""),
                    "location": item.get("location", ""),
                    "raw": item
                },
                timestamp=timestamp,
                author=item.get("owner"),
                url=f"https://app.quartzy.com/inventory/{item_id}" if item_id else None,
                doc_type="inventory_item"
            )

        except Exception as e:
            print(f"[Quartzy] Error converting inventory item: {e}")
            return None

    def _order_request_to_document(self, req: Dict[str, Any]) -> Optional[Document]:
        """Convert a Quartzy order request to Document"""
        try:
            req_id = str(req.get("id", ""))
            name = req.get("name", req.get("item_name", "Untitled Request"))

            parts = [f"Order Request: {name}"]

            if req.get("status"):
                parts.append(f"Status: {req['status']}")
            if req.get("vendor"):
                parts.append(f"Vendor: {req['vendor']}")
            if req.get("catalog_number"):
                parts.append(f"Catalog #: {req['catalog_number']}")
            if req.get("quantity"):
                parts.append(f"Quantity: {req['quantity']}")
            if req.get("unit_price") or req.get("price"):
                parts.append(f"Price: {req.get('unit_price', req.get('price', ''))}")
            if req.get("requester") or req.get("requested_by"):
                parts.append(f"Requested By: {req.get('requester', req.get('requested_by', ''))}")
            if req.get("lab"):
                parts.append(f"Lab: {req['lab']}")
            if req.get("notes"):
                parts.append(f"Notes: {req['notes']}")
            if req.get("url"):
                parts.append(f"Product URL: {req['url']}")

            content = "\n".join(parts)

            timestamp = None
            if req.get("created_at") or req.get("requested_at"):
                try:
                    ts = req.get("created_at") or req.get("requested_at")
                    timestamp = datetime.fromisoformat(str(ts).replace("Z", "+00:00"))
                except Exception:
                    pass

            return Document(
                doc_id=f"quartzy_order_{req_id}",
                source="quartzy",
                content=content,
                title=f"Order: {name}",
                metadata={
                    "quartzy_id": req_id,
                    "item_type": "order_request",
                    "status": req.get("status", ""),
                    "vendor": req.get("vendor", ""),
                    "raw": req
                },
                timestamp=timestamp,
                author=req.get("requester", req.get("requested_by")),
                url=f"https://app.quartzy.com/requests/{req_id}" if req_id else None,
                doc_type="order_request"
            )

        except Exception as e:
            print(f"[Quartzy] Error converting order request: {e}")
            return None

    @staticmethod
    def parse_csv(file_bytes: bytes, filename: str) -> List[Document]:
        """
        Parse a Quartzy CSV or Excel export into Document objects.

        Supports both .csv and .xlsx files exported from Quartzy.

        Expected columns (flexible matching):
        - Item Name, Serial Number, Vendor, Catalog #,
          Location, Owner, Unit Size, Amount in Stock, Type, etc.
        """
        documents = []

        try:
            ext = os.path.splitext(filename)[1].lower()

            if ext in ('.xlsx', '.xls'):
                # Parse Excel
                rows = QuartzyConnector._parse_excel(file_bytes)
            else:
                # Parse CSV (default)
                rows = QuartzyConnector._parse_csv_rows(file_bytes)

            print(f"[Quartzy CSV] Parsing {len(rows)} rows from {filename}")

            for idx, row in enumerate(rows):
                # Flexible column matching
                name = (
                    row.get("Item Name") or
                    row.get("item_name") or
                    row.get("Name") or
                    row.get("name") or
                    ""
                ).strip()

                if not name:
                    continue

                serial = row.get("Serial Number", row.get("serial_number", ""))
                vendor = row.get("Vendor", row.get("vendor", ""))
                catalog = row.get("Catalog #", row.get("catalog_number", row.get("Catalog Number", "")))
                location = row.get("Location", row.get("location", ""))
                owner = row.get("Owner", row.get("owner", ""))
                unit_size = row.get("Unit Size", row.get("unit_size", ""))
                stock = row.get("Amount in Stock", row.get("amount_in_stock", row.get("Quantity", "")))
                item_type = row.get("Type", row.get("type", ""))

                # Build content
                parts = [f"Lab Inventory Item: {name}"]
                if vendor:
                    parts.append(f"Vendor: {vendor}")
                if catalog:
                    parts.append(f"Catalog #: {catalog}")
                if location:
                    parts.append(f"Location: {location}")
                if owner:
                    parts.append(f"Owner: {owner}")
                if unit_size:
                    parts.append(f"Unit Size: {unit_size}")
                if stock:
                    parts.append(f"Amount in Stock: {stock}")
                if item_type:
                    parts.append(f"Type: {item_type}")

                # Include any extra columns
                known_cols = {
                    "Item Name", "item_name", "Name", "name",
                    "Serial Number", "serial_number",
                    "Vendor", "vendor",
                    "Catalog #", "catalog_number", "Catalog Number",
                    "Location", "location",
                    "Owner", "owner",
                    "Unit Size", "unit_size",
                    "Amount in Stock", "amount_in_stock", "Quantity",
                    "Type", "type",
                    "Row Type", "Delete", "Archived", "Instance ID"
                }
                for col, val in row.items():
                    if col not in known_cols and val and str(val).strip():
                        parts.append(f"{col}: {val}")

                content = "\n".join(parts)

                doc_id = f"quartzy_csv_{serial}" if serial else f"quartzy_csv_{idx}"

                documents.append(Document(
                    doc_id=doc_id,
                    source="quartzy_csv",
                    content=content,
                    title=f"Inventory: {name}",
                    metadata={
                        "serial_number": serial,
                        "vendor": vendor,
                        "catalog_number": catalog,
                        "location": location,
                        "owner": owner,
                        "source_file": filename,
                        "row_index": idx
                    },
                    timestamp=datetime.now(timezone.utc),
                    author=owner or None,
                    doc_type="inventory_item"
                ))

            print(f"[Quartzy CSV] Parsed {len(documents)} items from {filename}")

        except Exception as e:
            print(f"[Quartzy CSV] Parse error: {e}")
            import traceback
            traceback.print_exc()

        return documents

    @staticmethod
    def _parse_csv_rows(file_bytes: bytes) -> List[Dict[str, str]]:
        """Parse CSV bytes into list of row dicts"""
        rows = []
        try:
            # Try UTF-8 with BOM first
            text = file_bytes.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = file_bytes.decode("latin-1", errors="ignore")

        reader = csv.DictReader(io.StringIO(text))
        for row in reader:
            rows.append(dict(row))
        return rows

    @staticmethod
    def _parse_excel(file_bytes: bytes) -> List[Dict[str, str]]:
        """Parse Excel bytes into list of row dicts"""
        rows = []
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True)
            ws = wb.active
            if not ws:
                return rows

            headers = []
            for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
                if row_idx == 0:
                    headers = [str(cell) if cell else f"col_{i}" for i, cell in enumerate(row)]
                    continue
                row_dict = {}
                for i, cell in enumerate(row):
                    if i < len(headers):
                        row_dict[headers[i]] = str(cell) if cell is not None else ""
                rows.append(row_dict)

            wb.close()
        except ImportError:
            print("[Quartzy CSV] openpyxl not installed for Excel parsing")
        except Exception as e:
            print(f"[Quartzy CSV] Excel parse error: {e}")

        return rows
